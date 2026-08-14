import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from build_xlsx import build_workbook
from core.session_store import InMemorySessionStore
from modules.accident_analysis.presentation import AccidentPresentationService
from modules.accident_analysis.presentation.contract import (
    EXPECTED_TOKENS_BY_SLIDE,
    TRADITIONAL_SLIDE_SIZE,
    chart_frames,
    tokens_by_slide,
    validate_traditional_template,
)
from modules.accident_analysis.presentation.common import WIDE
from tests.test_accident_analysis import dataset
from web_server import content_disposition, parse_uploads, recover_major_violations
from tests.test_excel_sources import make_workbook


class SessionStoreTests(unittest.TestCase):
    def test_expiry_refresh_and_capacity_are_enforced(self):
        now = [100.0]
        store = InMemorySessionStore(ttl_seconds=10, max_sessions=1, clock=lambda: now[0])
        first = store.create({"name": "first"})
        now[0] = 105.0
        self.assertEqual(store.get(first)["name"], "first")
        now[0] = 106.0
        second = store.create({"name": "second"})
        self.assertIsNone(store.get(first))
        self.assertEqual(store.get(second)["name"], "second")
        now[0] = 117.0
        self.assertIsNone(store.get(second))


class RuntimeTests(unittest.TestCase):
    def test_multipart_upload_keeps_chinese_filename_and_binary_xlsx_bytes(self):
        boundary = b"----traffic-boundary"
        payload = b"PK\x03\x04\xe0\xffbinary"
        raw = b"\r\n".join((
            b"--" + boundary,
            'Content-Disposition: form-data; name="target"; filename="115年績效.xlsx"'.encode(),
            b"Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            b"",
            payload,
            b"--" + boundary + b"--",
            b"",
        ))
        handler = type("Handler", (), {"headers": {"Content-Length": str(len(raw)), "Content-Type": f"multipart/form-data; boundary={boundary.decode()}"}, "rfile": BytesIO(raw)})()
        upload = parse_uploads(handler, field_names={"target"})[0]
        self.assertEqual(upload.filename, "115年績效.xlsx")
        self.assertEqual(upload.file.read(), payload)

    def test_chinese_download_filename_uses_utf8_header(self):
        header = content_disposition("交通事故分析週報_新式版本.pptx")
        self.assertIn("filename=download.pptx", header)
        self.assertIn("filename*=UTF-8''", header)
        self.assertIn("%E4%BA%A4%E9%80%9A", header)

    def test_major_violation_batch_keeps_original_filenames(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            first = root / "第35條.xlsx"
            second = root / "第53條.xlsx"
            make_workbook(first, [("資料", [["違規日期", "違規條款"], [1150101, "35條"]])])
            make_workbook(second, [("資料", [["違規日期", "違規條款"], [1150102, "53條1項"]])])
            uploads = []
            for path in (first, second):
                upload = type("Upload", (), {})()
                upload.filename = path.name
                upload.file = BytesIO(path.read_bytes())
                uploads.append(upload)
            dataset = recover_major_violations(uploads)

        self.assertEqual(dataset.raw_row_count, 2)
        self.assertEqual([source.file_name for source in dataset.sources], ["第35條.xlsx", "第53條.xlsx"])


class PresentationServiceTests(unittest.TestCase):
    def test_traditional_template_matches_versioned_token_contract(self):
        template = AccidentPresentationService().variants["traditional"].template
        presentation = Presentation(template)
        validate_traditional_template(presentation)
        actual = tokens_by_slide(presentation)
        self.assertEqual(actual, EXPECTED_TOKENS_BY_SLIDE)
        self.assertEqual(len(set().union(*(counter.keys() for counter in actual))), 194)
        self.assertEqual(sum(sum(counter.values()) for counter in actual), 230)

    def test_traditional_template_validation_rejects_a_missing_token(self):
        template = AccidentPresentationService().variants["traditional"].template
        presentation = Presentation(template)
        period_shape = next(
            shape for shape in presentation.slides[0].shapes
            if getattr(shape, "has_text_frame", False) and "{{PERIOD}}" in shape.text
        )
        period_shape.text = period_shape.text.replace("{{PERIOD}}", "")
        with self.assertRaisesRegex(ValueError, "第 1 頁 TOKEN 不符"):
            validate_traditional_template(presentation)

    def test_unknown_variant_is_rejected_before_subprocess(self):
        with tempfile.TemporaryDirectory() as tempdir:
            service = AccidentPresentationService(Path(tempdir))
            with self.assertRaisesRegex(ValueError, "未知的投影片版本"):
                service.generate(dataset(), "unknown")

    def test_generation_creates_editable_pptx_without_external_runtime(self):
        service = AccidentPresentationService()
        for variant, filename in (("modern", "交通事故分析週報_新式版本.pptx"), ("traditional", "交通事故分析週報_傳統版本.pptx")):
            generated = service.generate(dataset(), variant, "115年1月1日至2月28日")
            self.assertTrue(generated.content.startswith(b"PK"))
            self.assertEqual(generated.filename, filename)
            presentation = Presentation(BytesIO(generated.content))
            self.assertEqual(len(presentation.slides), 7)
            expected_size = WIDE if variant == "modern" else TRADITIONAL_SLIDE_SIZE
            self.assertEqual((presentation.slide_width, presentation.slide_height), expected_size)
            picture_counts = [
                sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes)
                for slide in presentation.slides
            ]
            table_counts = [
                sum(shape.has_table for shape in slide.shapes)
                for slide in presentation.slides
            ]
            if variant == "modern":
                self.assertEqual(picture_counts, [0, 1, 1, 2, 2, 0, 0])
                self.assertEqual(table_counts, [0, 0, 0, 0, 0, 1, 0])
                layout_tolerance = Inches(0.1)
                for slide in presentation.slides:
                    for shape in slide.shapes:
                        self.assertGreaterEqual(shape.left, -layout_tolerance)
                        self.assertGreaterEqual(shape.top, -layout_tolerance)
                        self.assertLessEqual(shape.left + shape.width, presentation.slide_width + layout_tolerance)
                        self.assertLessEqual(shape.top + shape.height, presentation.slide_height + layout_tolerance)
            else:
                self.assertEqual(picture_counts, [1, 1, 1, 2, 2, 0, 0])
                self.assertEqual(table_counts, [0, 1, 1, 2, 2, 1, 5])
                template = Presentation(service.variants["traditional"].template)
                expected_frames = {frame[1:] for frame in chart_frames(template).values()}
                actual_pictures = {
                    (shape.left, shape.top, shape.width, shape.height)
                    for slide in presentation.slides for shape in slide.shapes
                    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
                }
                self.assertTrue(expected_frames.issubset(actual_pictures))
            with ZipFile(BytesIO(generated.content)) as archive:
                slide_xml = b"".join(
                    archive.read(name) for name in archive.namelist()
                    if name.startswith("ppt/slides/slide") and name.endswith(".xml")
                )
            self.assertNotIn(b"{{", slide_xml)


class ExcelExportTests(unittest.TestCase):
    def test_export_creates_formula_driven_workbook(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "analysis.xlsx"
            build_workbook({
                "headers": ["路段", "交叉路名", "發生時間", "肇事原因", "年齡", "當事者區分", "件數"],
                "rows": [["文化路", "中山路", "08:00", "未依規定讓車", 25, "機車", 3]],
            }, output)
            workbook = load_workbook(output, data_only=False)
        self.assertEqual(workbook.sheetnames, ["原始資料", "分析表"])
        self.assertEqual(workbook["原始資料"]["H2"].value[:3], "=IF")
        self.assertTrue(workbook["分析表"]["B6"].value.startswith("=SUMIF"))


if __name__ == "__main__":
    unittest.main()
