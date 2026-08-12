import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from pptx import Presentation

from modules.major_violation.data.transformer import transform_violation_source
from modules.major_violation.performance import (
    build_performance, generate_performance_pptx, load_performance_statistics, load_performance_targets,
)


class MajorViolationPerformanceTests(unittest.TestCase):
    def test_target_file_drives_unit_totals_and_presentation(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            target_path = root / "targets.xlsx"
            workbook = Workbook(); sheet = workbook.active; sheet.title = "13大項月"
            sheet.cell(4, 3, "闖紅燈"); sheet.cell(4, 4, "逆向行駛"); sheet.cell(4, 5, "轉彎未依規定")
            sheet.cell(6, 1, "板橋所"); sheet.cell(6, 2, "目標值"); sheet.cell(6, 3, 10); sheet.cell(6, 4, 4); sheet.cell(6, 5, 8)
            sheet.cell(9, 1, "後埔所"); sheet.cell(9, 2, "目標值"); sheet.cell(9, 3, 5); sheet.cell(9, 4, 2); sheet.cell(9, 5, 3)
            workbook.save(target_path)
            targets = load_performance_targets(target_path)
            dataset = transform_violation_source({"headers": ["違規日期", "舉發單位", "條次", "項次"], "rows": [[1150801, "板橋所", "53", "1"], [1150802, "板橋所", "53", "1"]]})
            result = build_performance(dataset, targets, "2026-08-01", "2026-08-09")
            self.assertEqual([item["label"] for item in result["categories"]], ["闖紅燈", "逆向行駛", "轉彎未依規定"])
            self.assertEqual(result["rows"][0]["cells"][0]["actual"], 2)
            output = root / "performance.pptx"
            template = Path(__file__).resolve().parents[1] / "modules" / "major_violation" / "templates" / "重大違規績效投影片.pptx"
            generate_performance_pptx(result, template, output)
            self.assertEqual(len(Presentation(BytesIO(output.read_bytes())).slides), 1)

    def test_statistics_file_overrides_actual_and_exposes_available_items(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "statistics.xlsx"
            workbook = Workbook(); sheet = workbook.active; sheet.title = "13大項月"
            sheet.cell(4, 3, "闖紅燈"); sheet.cell(6, 1, "板橋所"); sheet.cell(6, 2, "取締件數"); sheet.cell(6, 3, 18)
            workbook.save(path)
            statistics = load_performance_statistics([path])
            self.assertEqual(statistics.available_keys, ("red_light",))

    def test_weekly_target_uses_benchmark_sheet(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "weekly.xlsx"
            workbook = Workbook(); sheet = workbook.active; sheet.title = "基準值"
            sheet.cell(2, 2, "闖紅燈"); sheet.cell(5, 1, "板橋所"); sheet.cell(5, 2, 12)
            workbook.save(path)
            self.assertEqual(load_performance_targets(path, "week").values["板橋所"]["red_light"], 12)
