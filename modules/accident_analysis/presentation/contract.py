"""Versioned structural contract for the traditional accident PPTX template."""

from __future__ import annotations

import re
from collections import Counter

from pptx.util import Inches


TOKEN_PATTERN = re.compile(r"\{\{([A-Z0-9_]+)\}\}")
CHART_TOKENS = (
    "OVERVIEW_CHART", "ROAD_CHART", "INTERSECTION_CHART", "CAUSE_CHART",
    "TIME_CHART", "AGE_CHART", "VEHICLE_CHART",
)
TRADITIONAL_SLIDE_SIZE = (Inches(10), Inches(7.5))
TABLE_CONTRACTS = {
    "R1": (3, 11),
    "I1": (3, 11),
    "T1": (2, 13),
    "C1": (3, 13),
    "AAG1": (2, 8),
    "V1": (2, 8),
    "A1T": (9, 7),
}


def _tokens(*values: str) -> Counter:
    return Counter(values)


def _numbered(prefix: str, count: int) -> list[str]:
    return [f"{prefix}{index}" for index in range(1, count + 1)]


def _expected_tokens() -> tuple[Counter, ...]:
    slide1 = _tokens("PERIOD", "TOTAL", "TD", "A1", "A1D", "A2", "A2D", "SUMMARY", "OVERVIEW_CHART")
    slide2 = _tokens(*_numbered("R", 10), *_numbered("N", 10), "R1", "N1", "R2", "N2", "R3", "N3", "RCOUNT", "ROAD_CHART")
    slide3 = _tokens(*_numbered("I", 10), *_numbered("IN", 10), "I1", "IN1", "I2", "IN2", "I3", "IN3", "INTERSECTION_CHART")
    slide4 = _tokens(
        "C1", "CN1", "CP1", "C2", "CN2", "CP2", "C3", "CN3", "CP3",
        *_numbered("T", 12), *_numbered("TN", 12),
        "T1", "TN1", "TP1", "T2", "TN2", "TP2", "T3", "TN3", "TP3",
        *_numbered("C", 12), *_numbered("CN", 12), "CAUSE_CHART", "TIME_CHART",
    )
    slide5 = _tokens(
        *_numbered("AAG", 7), *_numbered("AN", 7),
        "AAG1", "AN1", "AP1", "AAG2", "AN2", "AP2", "AAG3", "AN3", "AP3",
        *_numbered("V", 7), *_numbered("VN", 7),
        "V1", "VN1", "VP1", "V2", "VN2", "VP2", "V3", "VN3", "VP3",
        "AGE_CHART", "VEHICLE_CHART",
    )
    details = [f"A{index}{suffix}" for index in range(1, 8) for suffix in "TLCVFN"]
    slide6 = _tokens("YEAR", *details)
    slide7 = _tokens("EPERIOD", "HOTSPOT", "FOCUS1", "FOCUS2", "AHOURS", "ACOUNT", "ALIMIT")
    return slide1, slide2, slide3, slide4, slide5, slide6, slide7


EXPECTED_TOKENS_BY_SLIDE = _expected_tokens()


def shape_text(shape) -> str:
    if shape.has_table:
        return "\n".join(cell.text for row in shape.table.rows for cell in row.cells)
    return shape.text if getattr(shape, "has_text_frame", False) else ""


def tokens_by_slide(presentation) -> tuple[Counter, ...]:
    return tuple(
        Counter(token for shape in slide.shapes for token in TOKEN_PATTERN.findall(shape_text(shape)))
        for slide in presentation.slides
    )


def find_table_by_token(presentation, token: str):
    marker = f"{{{{{token}}}}}"
    matches = [
        shape.table
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_table and marker in shape_text(shape)
    ]
    if len(matches) != 1:
        raise ValueError(f"傳統簡報模板的表格 TOKEN {marker} 應出現一次，實際為 {len(matches)} 次。")
    return matches[0]


def chart_frames(presentation) -> dict[str, tuple[object, int, int, int, int]]:
    frames = {}
    for token in CHART_TOKENS:
        marker = f"{{{{{token}}}}}"
        matches = [
            (slide, shape)
            for slide in presentation.slides
            for shape in slide.shapes
            if not shape.has_table and marker in shape_text(shape)
        ]
        if len(matches) != 1:
            raise ValueError(f"傳統簡報模板的圖表 TOKEN {marker} 應出現一次，實際為 {len(matches)} 次。")
        slide, shape = matches[0]
        frames[token] = (slide, shape.left, shape.top, shape.width, shape.height)
    return frames


def validate_traditional_template(presentation) -> None:
    errors = []
    if len(presentation.slides) != len(EXPECTED_TOKENS_BY_SLIDE):
        errors.append(f"投影片應為 7 頁，實際為 {len(presentation.slides)} 頁")
    if (presentation.slide_width, presentation.slide_height) != TRADITIONAL_SLIDE_SIZE:
        errors.append("投影片尺寸應為 10 × 7.5 吋（4:3）")
    actual_by_slide = tokens_by_slide(presentation)
    for index, expected in enumerate(EXPECTED_TOKENS_BY_SLIDE, 1):
        actual = actual_by_slide[index - 1] if index <= len(actual_by_slide) else Counter()
        if actual != expected:
            missing = list((expected - actual).elements())
            unknown = list((actual - expected).elements())
            errors.append(f"第 {index} 頁 TOKEN 不符：缺少 {missing or '無'}；多出 {unknown or '無'}")
    if not errors:
        for token, expected_size in TABLE_CONTRACTS.items():
            table = find_table_by_token(presentation, token)
            actual_size = (len(table.rows), len(table.columns))
            if actual_size != expected_size:
                errors.append(f"TOKEN {{{{{token}}}}} 所在表格應為 {expected_size[0]}×{expected_size[1]}，實際為 {actual_size[0]}×{actual_size[1]}")
        chart_frames(presentation)
    if errors:
        raise ValueError("傳統簡報模板契約不符：" + "；".join(errors))
