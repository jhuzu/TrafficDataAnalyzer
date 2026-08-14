"""Template-filled 4:3 accident presentation."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from .common import chart, remove_shape, replace_tokens_in_text_frame
from .contract import CHART_TOKENS, find_table_by_token, shape_text, validate_traditional_template


def _traditional_replacements(payload: dict) -> dict[str, object]:
    overview, rankings, distributions = payload["overview"], payload["rankings"], payload["distributions"]
    fmt = lambda value: f"{float(value or 0):,.0f}"
    pct = lambda item: f'{float(item.get("share", 0)) * 100:.1f}%'
    replacements = {
        "PERIOD": payload["period"]["display"], "YEAR": payload["period"]["year"],
        "SUMMARY": payload["narrative"]["summary"], "TOTAL": fmt(overview["total"]),
        "TD": overview["changes"]["total"]["text"], "A1": fmt(overview["a1"]),
        "A1D": overview["changes"]["a1"]["text"], "A2": fmt(overview["a2"]),
        "A2D": overview["changes"]["a2"]["text"], "RCOUNT": fmt(overview["roadCategoryCount"]),
        "EPERIOD": payload["period"]["display"],
        "HOTSPOT": rankings["roads"][0]["label"] if rankings["roads"] else "主要事故熱點",
        "AHOURS": max(1, round(float(overview["total"]) / 200)),
        "ACOUNT": max(1, round(float(overview["total"]) / 20)), "ALIMIT": 3,
    }
    for prefix, items in (
        ("R", rankings["roads"]), ("I", rankings["intersections"]),
        ("C", rankings["causes"]), ("T", distributions["timeBuckets"]),
        ("AAG", distributions["ageGroups"]), ("V", rankings["vehicles"]),
    ):
        count_prefix = {"R": "N", "I": "IN", "C": "CN", "T": "TN", "AAG": "AN", "V": "VN"}[prefix]
        for index in range(1, 13):
            item = items[index - 1] if index <= len(items) else {}
            replacements[f"{prefix}{index}"] = item.get("label", "")
            replacements[f"{count_prefix}{index}"] = fmt(item.get("count", 0)) if item else ""
    for prefix, items in (
        ("CP", rankings["causes"]), ("TP", distributions["rankedTimeBuckets"]),
        ("AP", distributions["ageGroups"]), ("VP", rankings["vehicles"]),
    ):
        for index in range(1, 4):
            replacements[f"{prefix}{index}"] = pct(items[index - 1]) if index <= len(items) else ""
    a1_roads = rankings["a1Roads"] or rankings["roads"]
    replacements["FOCUS1"] = a1_roads[0]["label"] if a1_roads else "主要事故熱點"
    replacements["FOCUS2"] = a1_roads[1]["label"] if len(a1_roads) > 1 else replacements["FOCUS1"]
    for index in range(1, 8):
        replacements.update({f"A{index}T": "", f"A{index}L": "", f"A{index}C": "", f"A{index}V": "", f"A{index}F": "", f"A{index}N": ""})
    for index, item in enumerate(payload["a1Details"][:7], 1):
        replacements.update({
            f"A{index}T": item["time"], f"A{index}L": item["location"],
            f"A{index}C": item["cause"], f"A{index}V": item["vehicle"],
            f"A{index}F": f'{item["deaths"]}人', f"A{index}N": item["notes"],
        })
    return replacements


def _fill_template_tables(presentation, payload: dict) -> None:
    rankings, distributions = payload["rankings"], payload["distributions"]

    def put(table, row, column, value):
        table.cell(row, column).text = str(value)

    for token, items in (("R1", rankings["roads"][:10]), ("I1", rankings["intersections"][:10])):
        table = find_table_by_token(presentation, token)
        for index, item in enumerate(items, 1):
            put(table, 0, index, item["label"])
            put(table, 1, index, f'{item["count"]:,.0f}')
            put(table, 2, index, index)
    time_table = find_table_by_token(presentation, "T1")
    for index, item in enumerate(distributions["timeBuckets"][:12], 1):
        put(time_table, 0, index, item["label"])
        put(time_table, 1, index, f'{item["count"]:,.0f}')
    cause_table = find_table_by_token(presentation, "C1")
    for index, item in enumerate(rankings["causes"][:12], 1):
        put(cause_table, 1, index, item["label"])
        put(cause_table, 2, index, f'{item["count"]:,.0f}')
    age_table = find_table_by_token(presentation, "AAG1")
    for index, item in enumerate(distributions["ageGroups"][:7], 1):
        put(age_table, 0, index, item["label"])
        put(age_table, 1, index, f'{item["count"]:,.0f}')
    vehicle_table = find_table_by_token(presentation, "V1")
    for index, item in enumerate(rankings["vehicles"][:7], 1):
        put(vehicle_table, 0, index, item["label"])
        put(vehicle_table, 1, index, f'{item["count"]:,.0f}')
    detail_table = find_table_by_token(presentation, "A1T")
    for row, item in enumerate(payload["a1Details"][:7], 2):
        for column, value in enumerate((item["number"], item["time"], item["location"], item["cause"], item["vehicle"], f'{item["deaths"]}人', item["notes"])):
            put(detail_table, row, column, value)


def _replace_tokens(presentation, replacements: dict[str, object]):
    chart_frames = {}
    for slide in presentation.slides:
        for shape in list(slide.shapes):
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        replace_tokens_in_text_frame(cell.text_frame, replacements)
                continue
            if not getattr(shape, "has_text_frame", False):
                continue
            matched = next((token for token in CHART_TOKENS if f"{{{{{token}}}}}" in shape_text(shape)), None)
            if matched:
                chart_frames[matched] = (slide, shape.left, shape.top, shape.width, shape.height)
                remove_shape(shape)
            else:
                replace_tokens_in_text_frame(shape.text_frame, replacements)
    return chart_frames


def generate_traditional_presentation(payload: dict, template_path: str | Path, output_path: str | Path) -> None:
    """Fill the established 4:3 operational template without changing its layout."""
    presentation = Presentation(template_path)
    validate_traditional_template(presentation)
    _fill_template_tables(presentation, payload)
    frames = _replace_tokens(presentation, _traditional_replacements(payload))
    overview, rankings, distributions = payload["overview"], payload["rankings"], payload["distributions"]
    chart_data = {
        "OVERVIEW_CHART": ([{"label": "A1", "count": overview["a1"]}, {"label": "A2", "count": overview["a2"]}], "C00000"),
        "ROAD_CHART": (rankings["roads"][:10], "2F75B5"),
        "INTERSECTION_CHART": (rankings["intersections"][:10], "ED7D31"),
        "CAUSE_CHART": (rankings["causes"][:8], "5B9BD5"),
        "TIME_CHART": (distributions["timeBuckets"], "70AD47"),
        "AGE_CHART": (distributions["ageGroups"][:7], "4472C4"),
        "VEHICLE_CHART": (rankings["vehicles"][:7], "ED7D31"),
    }
    for token, (items, color) in chart_data.items():
        slide, left, top, width, height = frames[token]
        slide.shapes.add_picture(chart(items, color=color), left, top, width=width, height=height)
    presentation.save(output_path)
