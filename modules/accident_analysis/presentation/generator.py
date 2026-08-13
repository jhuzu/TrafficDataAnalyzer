"""Public-Python PPTX generator for accident analysis reports."""

from __future__ import annotations

import platform
from io import BytesIO
from pathlib import Path
from typing import Iterable

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


WIDE = (Inches(13.333), Inches(7.5))
THEMES = {
    "modern": {"accent": "1F4E78", "accent2": "ED7D31", "pale": "EAF2F8", "name": "新式版本"},
    "traditional": {"accent": "305496", "accent2": "C00000", "pale": "F3F6FB", "name": "傳統版本"},
}
TEXT_FONT_NAME = "Microsoft JhengHei" if platform.system() == "Windows" else "PingFang TC"


def _font(size: int, bold: bool = False):
    candidates = (
        ("C:/Windows/Fonts/msjhbd.ttc", "C:/Windows/Fonts/msjh.ttc")
        if platform.system() == "Windows" and bold
        else ("C:/Windows/Fonts/msjh.ttc", "C:/Windows/Fonts/msjhbd.ttc")
        if platform.system() == "Windows"
        else ("/System/Library/Fonts/PingFang.ttc", "/System/Library/Fonts/STHeiti Light.ttc")
    )
    for candidate in candidates:
        if Path(candidate).is_file():
            return ImageFont.truetype(candidate, size, index=0)
    return ImageFont.load_default()


def _chart(items: Iterable[dict], width=1100, height=470, color="#1F4E78") -> BytesIO:
    values = list(items) or [{"label": "無資料", "count": 0}]
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font, label_font, value_font = _font(26, True), _font(18), _font(18, True)
    maximum = max(1, *(float(item.get("count", 0)) for item in values))
    usable_height = height - 56
    row_height = max(34, usable_height // len(values))
    for index, item in enumerate(values):
        y = 30 + index * row_height
        label = str(item.get("label", ""))[:22]
        amount = float(item.get("count", 0))
        draw.text((18, y), label, fill="#26364A", font=label_font)
        left, bar_y = 350, y + 4
        bar_width = int((width - 450) * amount / maximum)
        draw.rounded_rectangle((left, bar_y, left + bar_width, bar_y + 22), 8, fill=f"#{color.lstrip('#')}")
        draw.text((min(width - 85, left + bar_width + 10), y), f"{amount:,.0f}", fill="#26364A", font=value_font)
    output = BytesIO()
    image.save(output, format="PNG")
    output.seek(0)
    return output


def _set_background(slide, theme):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb("FFFFFF")
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, WIDE[0], Inches(0.22))
    bar.fill.solid(); bar.fill.fore_color.rgb = _rgb(theme["accent"]); bar.line.fill.background()


def _rgb(value):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(value)


def _text(slide, text, left, top, width, height, size=18, bold=False, color="26364A", align=None):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear(); frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = str(text)
    paragraph.font.name = TEXT_FONT_NAME
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = _rgb(color)
    if align is not None:
        paragraph.alignment = align
    return box


def _title(slide, title, subtitle, theme):
    _set_background(slide, theme)
    _text(slide, title, 0.65, 0.42, 12, 0.48, 29, True, theme["accent"])
    _text(slide, subtitle, 0.67, 0.95, 12, 0.32, 12, False, "52606D")


def _card(slide, label, value, detail, left, theme):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(1.55), Inches(3.75), Inches(1.55))
    shape.fill.solid(); shape.fill.fore_color.rgb = _rgb(theme["pale"]); shape.line.color.rgb = _rgb("C9D7E5")
    _text(slide, label, left + 0.24, 1.8, 3.2, 0.3, 14, True, theme["accent"])
    _text(slide, value, left + 0.24, 2.12, 3.2, 0.45, 25, True, theme["accent2"])
    _text(slide, detail, left + 0.24, 2.65, 3.2, 0.26, 11, False, "52606D")


def _table(slide, columns, rows, left=0.7, top=1.5, width=12.0, height=4.9, theme=None):
    table = slide.shapes.add_table(len(rows) + 1, len(columns), Inches(left), Inches(top), Inches(width), Inches(height)).table
    for index, label in enumerate(columns):
        cell = table.cell(0, index); cell.text = label
        cell.fill.solid(); cell.fill.fore_color.rgb = _rgb(theme["accent"])
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.size = Pt(12); paragraph.font.bold = True; paragraph.font.color.rgb = _rgb("FFFFFF")
    for row_index, row in enumerate(rows, 1):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index); cell.text = str(value)
            cell.fill.solid(); cell.fill.fore_color.rgb = _rgb("FFFFFF" if row_index % 2 else theme["pale"])
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(10); paragraph.font.color.rgb = _rgb("26364A")
    return table


def _rank_rows(items, total, limit=10):
    return [[index, item.get("label", ""), f'{float(item.get("count", 0)):,.0f}', f'{float(item.get("share", 0)) * 100:.1f}%'] for index, item in enumerate(items[:limit], 1)]


def _remove_shape(shape):
    element = shape._element
    element.getparent().remove(element)


def _replace_tokens(presentation, replacements):
    chart_frames = {}
    for slide in presentation.slides:
        for shape in list(slide.shapes):
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        text = cell.text
                        for token, value in replacements.items():
                            text = text.replace(f"{{{{{token}}}}}", str(value))
                        if text != cell.text:
                            cell.text = text
                continue
            if not getattr(shape, "has_text_frame", False):
                continue
            text = shape.text
            for token in ("OVERVIEW_CHART", "ROAD_CHART", "INTERSECTION_CHART", "CAUSE_CHART", "TIME_CHART", "AGE_CHART", "VEHICLE_CHART"):
                if f"{{{{{token}}}}}" in text:
                    chart_frames[token] = (slide, shape.left, shape.top, shape.width, shape.height)
                    _remove_shape(shape)
                    break
            else:
                for token, value in replacements.items():
                    text = text.replace(f"{{{{{token}}}}}", str(value))
                if text != shape.text:
                    shape.text = text
    return chart_frames


def _fill_template_tables(presentation, payload):
    rankings, distributions = payload["rankings"], payload["distributions"]
    by_slide = {}
    for number, slide in enumerate(presentation.slides, 1):
        by_slide[number] = [shape.table for shape in slide.shapes if shape.has_table]
    def put(table, row, column, value):
        if row < len(table.rows) and column < len(table.columns):
            table.cell(row, column).text = str(value)
    for slide_number, items in ((2, rankings["roads"][:10]), (3, rankings["intersections"][:10])):
        if by_slide[slide_number]:
            for index, item in enumerate(items, 1):
                put(by_slide[slide_number][0], 0, index, item["label"]); put(by_slide[slide_number][0], 1, index, f'{item["count"]:,.0f}'); put(by_slide[slide_number][0], 2, index, index)
    if len(by_slide[4]) >= 2:
        for index, item in enumerate(distributions["timeBuckets"][:12], 1):
            put(by_slide[4][0], 0, index, item["label"]); put(by_slide[4][0], 1, index, f'{item["count"]:,.0f}')
        for index, item in enumerate(rankings["causes"][:12], 1):
            put(by_slide[4][1], 1, index, item["label"]); put(by_slide[4][1], 2, index, f'{item["count"]:,.0f}')
    if len(by_slide[5]) >= 2:
        for index, item in enumerate(distributions["ageGroups"][:7], 1):
            put(by_slide[5][0], 0, index, item["label"]); put(by_slide[5][0], 1, index, f'{item["count"]:,.0f}')
        for index, item in enumerate(rankings["vehicles"][:7], 1):
            put(by_slide[5][1], 0, index, item["label"]); put(by_slide[5][1], 1, index, f'{item["count"]:,.0f}')
    if by_slide[6]:
        for row, item in enumerate(payload["a1Details"], 2):
            for column, value in enumerate((item["number"], item["time"], item["location"], item["cause"], item["vehicle"], f'{item["deaths"]}人', item["notes"])):
                put(by_slide[6][0], row, column, value)


def generate_traditional_presentation(payload: dict, template_path: str | Path, output_path: str | Path) -> None:
    """Fill the established 4:3 operational template without changing its layout."""
    presentation = Presentation(template_path)
    overview, rankings, distributions = payload["overview"], payload["rankings"], payload["distributions"]
    fmt = lambda value: f"{float(value or 0):,.0f}"
    pct = lambda item: f'{float(item.get("share", 0)) * 100:.1f}%'
    replacements = {"PERIOD": payload["period"]["display"], "YEAR": payload["period"]["year"], "SUMMARY": payload["narrative"]["summary"], "TOTAL": fmt(overview["total"]), "TD": overview["changes"]["total"]["text"], "A1": fmt(overview["a1"]), "A1D": overview["changes"]["a1"]["text"], "A2": fmt(overview["a2"]), "A2D": overview["changes"]["a2"]["text"], "RCOUNT": fmt(overview["roadCategoryCount"]), "EPERIOD": payload["period"]["display"], "HOTSPOT": rankings["roads"][0]["label"] if rankings["roads"] else "主要事故熱點", "AHOURS": max(1, round(float(overview["total"]) / 200)), "ACOUNT": max(1, round(float(overview["total"]) / 20)), "ALIMIT": 3}
    for prefix, items in (("R", rankings["roads"]), ("I", rankings["intersections"]), ("C", rankings["causes"]), ("T", distributions["timeBuckets"]), ("AAG", distributions["ageGroups"]), ("V", rankings["vehicles"])):
        count_prefix = {"R": "N", "I": "IN", "C": "CN", "T": "TN", "AAG": "AN", "V": "VN"}[prefix]
        for index in range(1, 13):
            item = items[index - 1] if index <= len(items) else {}
            replacements[f"{prefix}{index}"] = item.get("label", "")
            replacements[f"{count_prefix}{index}"] = fmt(item.get("count", 0)) if item else ""
    for prefix, items in (("CP", rankings["causes"]), ("TP", distributions["rankedTimeBuckets"]), ("AP", distributions["ageGroups"]), ("VP", rankings["vehicles"])):
        for index in range(1, 4): replacements[f"{prefix}{index}"] = pct(items[index - 1]) if index <= len(items) else ""
    a1_roads = rankings["a1Roads"] or rankings["roads"]
    replacements["FOCUS1"] = a1_roads[0]["label"] if a1_roads else "主要事故熱點"
    replacements["FOCUS2"] = a1_roads[1]["label"] if len(a1_roads) > 1 else replacements["FOCUS1"]
    for index in range(1, 8):
        replacements.update({f"A{index}T": "", f"A{index}L": "", f"A{index}C": "", f"A{index}V": "", f"A{index}F": "", f"A{index}N": ""})
    for index, item in enumerate(payload["a1Details"], 1):
        replacements.update({f"A{index}T": item["time"], f"A{index}L": item["location"], f"A{index}C": item["cause"], f"A{index}V": item["vehicle"], f"A{index}F": f'{item["deaths"]}人', f"A{index}N": item["notes"]})
    frames = _replace_tokens(presentation, replacements)
    _fill_template_tables(presentation, payload)
    chart_data = {"OVERVIEW_CHART": ([{"label": "A1", "count": overview["a1"]}, {"label": "A2", "count": overview["a2"]}], "C00000"), "ROAD_CHART": (rankings["roads"][:10], "2F75B5"), "INTERSECTION_CHART": (rankings["intersections"][:10], "ED7D31"), "CAUSE_CHART": (rankings["causes"][:8], "5B9BD5"), "TIME_CHART": (distributions["timeBuckets"], "70AD47"), "AGE_CHART": (distributions["ageGroups"][:7], "4472C4"), "VEHICLE_CHART": (rankings["vehicles"][:7], "ED7D31")}
    for token, (items, color) in chart_data.items():
        if token in frames:
            slide, left, top, width, height = frames[token]
            slide.shapes.add_picture(_chart(items, color=color), left, top, width=width, height=height)
    presentation.save(output_path)


def generate_presentation(payload: dict, variant_key: str, output_path: str | Path) -> None:
    """Generate a stable, editable PPTX from the existing JSON payload contract."""
    theme = THEMES[variant_key]
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = WIDE
    blank = presentation.slide_layouts[6]
    period = payload["period"]["display"]
    overview, rankings, distributions = payload["overview"], payload["rankings"], payload["distributions"]
    total, a1, a2 = overview["total"], overview["a1"], overview["a2"]

    slide = presentation.slides.add_slide(blank)
    _title(slide, "交通事故分析週報", f"{period}｜{theme['name']}", theme)
    _card(slide, "A1 + A2 事故總件數", f"{total:,.0f} 件", overview["changes"]["total"]["text"], 0.7, theme)
    _card(slide, "A1 事故", f"{a1:,.0f} 件", overview["changes"]["a1"]["text"], 4.78, theme)
    _card(slide, "A2 事故", f"{a2:,.0f} 件", overview["changes"]["a2"]["text"], 8.86, theme)
    _text(slide, payload["narrative"]["summary"], 0.9, 3.55, 11.6, 0.9, 17, False)
    _text(slide, "勤務建議", 0.9, 4.7, 3, 0.35, 18, True, theme["accent"])
    _text(slide, payload["narrative"]["enforcement"], 0.9, 5.12, 11.5, 0.9, 16)

    for title, key, color in (("易肇事路段排行", "roads", theme["accent"]), ("易肇事路口排行", "intersections", theme["accent2"])):
        slide = presentation.slides.add_slide(blank)
        _title(slide, title, period, theme)
        stream = _chart(rankings[key][:10], color=color)
        slide.shapes.add_picture(stream, Inches(0.7), Inches(1.35), width=Inches(12), height=Inches(5.45))

    slide = presentation.slides.add_slide(blank)
    _title(slide, "主要肇因與事故時段", period, theme)
    slide.shapes.add_picture(_chart(rankings["causes"][:7], width=540, height=410, color=theme["accent"]), Inches(0.6), Inches(1.35), width=Inches(6.0), height=Inches(4.6))
    slide.shapes.add_picture(_chart(distributions["rankedTimeBuckets"][:7], width=540, height=410, color=theme["accent2"]), Inches(6.75), Inches(1.35), width=Inches(5.95), height=Inches(4.6))
    _text(slide, payload["narrative"]["coreFocus"], 0.8, 6.25, 11.8, 0.35, 15, True, theme["accent"])

    slide = presentation.slides.add_slide(blank)
    _title(slide, "年齡與車種風險分布", period, theme)
    slide.shapes.add_picture(_chart(distributions["ageGroups"][:7], width=540, height=410, color=theme["accent"]), Inches(0.6), Inches(1.35), width=Inches(6.0), height=Inches(4.6))
    slide.shapes.add_picture(_chart(rankings["vehicles"][:7], width=540, height=410, color=theme["accent2"]), Inches(6.75), Inches(1.35), width=Inches(5.95), height=Inches(4.6))

    slide = presentation.slides.add_slide(blank)
    _title(slide, "A1 事故案件明細", period, theme)
    rows = [[item["number"], item["time"], item["location"], item["cause"], item["vehicle"], f'{item["deaths"]}人'] for item in payload["a1Details"]]
    _table(slide, ("編號", "發生時間", "發生地點", "肇因研判", "車種別", "死亡"), rows or [["", "", "無 A1 明細", "", "", ""]], theme=theme)

    slide = presentation.slides.add_slide(blank)
    _title(slide, "結論與防制作為", period, theme)
    _text(slide, "分析結論", 0.85, 1.55, 3, 0.4, 20, True, theme["accent"])
    _text(slide, payload["narrative"]["summary"], 0.9, 2.05, 11.4, 1.05, 19)
    _text(slide, "建議勤務部署", 0.85, 3.45, 3, 0.4, 20, True, theme["accent"])
    _text(slide, payload["narrative"]["enforcement"], 0.9, 3.95, 11.4, 1.05, 19)
    _text(slide, "統計口徑：以 Excel「件數」欄加總；本檔案由本機公開 Python 套件離線生成。", 0.9, 6.4, 11.5, 0.3, 11, False, "52606D")

    presentation.save(output_path)
