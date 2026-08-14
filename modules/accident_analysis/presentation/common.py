"""Low-level PPTX helpers shared by accident presentation builders."""

from __future__ import annotations

import platform
from io import BytesIO
from pathlib import Path
from typing import Iterable, Mapping

from PIL import Image, ImageDraw, ImageFont
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


WIDE = (Inches(13.333), Inches(7.5))
THEMES = {
    "modern": {"accent": "1F4E78", "accent2": "ED7D31", "pale": "EAF2F8", "name": "新式版本"},
    "traditional": {"accent": "305496", "accent2": "C00000", "pale": "F3F6FB", "name": "傳統版本"},
}
TEXT_FONT_NAME = "Microsoft JhengHei" if platform.system() == "Windows" else "PingFang TC"


def font(size: int, bold: bool = False):
    candidates = (
        ("C:/Windows/Fonts/msjhbd.ttc", "C:/Windows/Fonts/msjh.ttc")
        if platform.system() == "Windows" and bold
        else ("C:/Windows/Fonts/msjh.ttc", "C:/Windows/Fonts/msjhbd.ttc")
        if platform.system() == "Windows"
        else ("/System/Library/Fonts/PingFang.ttc", "/System/Library/Fonts/STHeiti Light.ttc")
    )
    for candidate in candidates:
        if Path(candidate).is_file():
            return ImageFont.truetype(candidate, size, index=0)
    return ImageFont.load_default()


def chart(items: Iterable[dict], width=1100, height=470, color="#1F4E78") -> BytesIO:
    values = list(items) or [{"label": "無資料", "count": 0}]
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    label_font, value_font = font(18), font(18, True)
    maximum = max(1, *(float(item.get("count", 0)) for item in values))
    usable_height = height - 56
    row_height = max(34, usable_height // len(values))
    for index, item in enumerate(values):
        y = 30 + index * row_height
        label = str(item.get("label", ""))[:22]
        amount = float(item.get("count", 0))
        draw.text((18, y), label, fill="#26364A", font=label_font)
        left, bar_y = 350, y + 4
        bar_width = int((width - 450) * amount / maximum)
        draw.rounded_rectangle((left, bar_y, left + bar_width, bar_y + 22), 8, fill=f"#{color.lstrip('#')}")
        draw.text((min(width - 85, left + bar_width + 10), y), f"{amount:,.0f}", fill="#26364A", font=value_font)
    output = BytesIO()
    image.save(output, format="PNG")
    output.seek(0)
    return output


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_background(slide, theme) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb("FFFFFF")
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, WIDE[0], Inches(0.22))
    bar.fill.solid()
    bar.fill.fore_color.rgb = rgb(theme["accent"])
    bar.line.fill.background()


def add_text(slide, value, left, top, width, height, size=18, bold=False, color="26364A", align=None):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = str(value)
    paragraph.font.name = TEXT_FONT_NAME
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = rgb(color)
    if align is not None:
        paragraph.alignment = align
    return box


def add_title(slide, title_value, subtitle, theme) -> None:
    set_background(slide, theme)
    add_text(slide, title_value, 0.65, 0.42, 12, 0.48, 29, True, theme["accent"])
    add_text(slide, subtitle, 0.67, 0.95, 12, 0.32, 12, False, "52606D")


def add_table(slide, columns, rows, left=0.7, top=1.5, width=12.0, height=4.9, theme=None):
    table = slide.shapes.add_table(
        len(rows) + 1, len(columns), Inches(left), Inches(top), Inches(width), Inches(height)
    ).table
    for index, label in enumerate(columns):
        cell = table.cell(0, index)
        cell.text = label
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(theme["accent"])
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.size = Pt(12)
            paragraph.font.bold = True
            paragraph.font.color.rgb = rgb("FFFFFF")
    for row_index, row in enumerate(rows, 1):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.text = str(value)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb("FFFFFF" if row_index % 2 else theme["pale"])
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(10)
                paragraph.font.color.rgb = rgb("26364A")
    return table


def remove_shape(shape) -> None:
    """Remove a shape; python-pptx currently has no public removal API."""
    element = shape._element
    element.getparent().remove(element)


def replace_tokens_in_text_frame(text_frame, replacements: Mapping[str, object]) -> None:
    """Replace tokens while preserving run formatting whenever token boundaries allow it."""
    for paragraph in text_frame.paragraphs:
        original = paragraph.text
        updated = original
        for token, value in replacements.items():
            updated = updated.replace(f"{{{{{token}}}}}", str(value))
        if updated == original:
            continue
        for run in paragraph.runs:
            run_value = run.text
            for token, value in replacements.items():
                run_value = run_value.replace(f"{{{{{token}}}}}", str(value))
            run.text = run_value
        if "".join(run.text for run in paragraph.runs) != updated:
            paragraph.text = updated
