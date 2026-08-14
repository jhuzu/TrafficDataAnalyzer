"""Code-built 16:9 accident presentation."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from .common import THEMES, WIDE, add_table, add_text, add_title, chart, rgb


def _card(slide, label, value, detail, left, theme) -> None:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(1.55), Inches(3.75), Inches(1.55)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(theme["pale"])
    shape.line.color.rgb = rgb("C9D7E5")
    add_text(slide, label, left + 0.24, 1.8, 3.2, 0.3, 14, True, theme["accent"])
    add_text(slide, value, left + 0.24, 2.12, 3.2, 0.45, 25, True, theme["accent2"])
    add_text(slide, detail, left + 0.24, 2.65, 3.2, 0.26, 11, False, "52606D")


def generate_presentation(payload: dict, variant_key: str, output_path: str | Path) -> None:
    """Generate a stable, editable PPTX from the shared JSON payload contract."""
    theme = THEMES[variant_key]
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = WIDE
    blank = presentation.slide_layouts[6]
    period = payload["period"]["display"]
    overview, rankings, distributions = payload["overview"], payload["rankings"], payload["distributions"]
    total, a1, a2 = overview["total"], overview["a1"], overview["a2"]

    slide = presentation.slides.add_slide(blank)
    add_title(slide, "交通事故分析週報", f"{period}｜{theme['name']}", theme)
    _card(slide, "A1 + A2 事故總件數", f"{total:,.0f} 件", overview["changes"]["total"]["text"], 0.7, theme)
    _card(slide, "A1 事故", f"{a1:,.0f} 件", overview["changes"]["a1"]["text"], 4.78, theme)
    _card(slide, "A2 事故", f"{a2:,.0f} 件", overview["changes"]["a2"]["text"], 8.86, theme)
    add_text(slide, payload["narrative"]["summary"], 0.9, 3.55, 11.6, 0.9, 17, False)
    add_text(slide, "勤務建議", 0.9, 4.7, 3, 0.35, 18, True, theme["accent"])
    add_text(slide, payload["narrative"]["enforcement"], 0.9, 5.12, 11.5, 0.9, 16)

    for title, key, color in (
        ("易肇事路段排行", "roads", theme["accent"]),
        ("易肇事路口排行", "intersections", theme["accent2"]),
    ):
        slide = presentation.slides.add_slide(blank)
        add_title(slide, title, period, theme)
        slide.shapes.add_picture(
            chart(rankings[key][:10], color=color), Inches(0.7), Inches(1.35),
            width=Inches(12), height=Inches(5.45),
        )

    slide = presentation.slides.add_slide(blank)
    add_title(slide, "主要肇因與事故時段", period, theme)
    slide.shapes.add_picture(chart(rankings["causes"][:7], width=540, height=410, color=theme["accent"]), Inches(0.6), Inches(1.35), width=Inches(6.0), height=Inches(4.6))
    slide.shapes.add_picture(chart(distributions["rankedTimeBuckets"][:7], width=540, height=410, color=theme["accent2"]), Inches(6.75), Inches(1.35), width=Inches(5.95), height=Inches(4.6))
    add_text(slide, payload["narrative"]["coreFocus"], 0.8, 6.25, 11.8, 0.35, 15, True, theme["accent"])

    slide = presentation.slides.add_slide(blank)
    add_title(slide, "年齡與車種風險分布", period, theme)
    slide.shapes.add_picture(chart(distributions["ageGroups"][:7], width=540, height=410, color=theme["accent"]), Inches(0.6), Inches(1.35), width=Inches(6.0), height=Inches(4.6))
    slide.shapes.add_picture(chart(rankings["vehicles"][:7], width=540, height=410, color=theme["accent2"]), Inches(6.75), Inches(1.35), width=Inches(5.95), height=Inches(4.6))

    slide = presentation.slides.add_slide(blank)
    add_title(slide, "A1 事故案件明細", period, theme)
    rows = [[item["number"], item["time"], item["location"], item["cause"], item["vehicle"], f'{item["deaths"]}人'] for item in payload["a1Details"]]
    add_table(slide, ("編號", "發生時間", "發生地點", "肇因研判", "車種別", "死亡"), rows or [["", "", "無 A1 明細", "", "", ""]], theme=theme)

    slide = presentation.slides.add_slide(blank)
    add_title(slide, "結論與防制作為", period, theme)
    add_text(slide, "分析結論", 0.85, 1.55, 3, 0.4, 20, True, theme["accent"])
    add_text(slide, payload["narrative"]["summary"], 0.9, 2.05, 11.4, 1.05, 19)
    add_text(slide, "建議勤務部署", 0.85, 3.45, 3, 0.4, 20, True, theme["accent"])
    add_text(slide, payload["narrative"]["enforcement"], 0.9, 3.95, 11.4, 1.05, 19)
    add_text(slide, "統計口徑：以 Excel「件數」欄加總；本檔案由本機公開 Python 套件離線生成。", 0.9, 6.4, 11.5, 0.3, 11, False, "52606D")
    presentation.save(output_path)
