"""Target-driven weekly performance table for major-violation reporting."""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .data.schema import ViolationDataset


PERFORMANCE_CATEGORIES = (
    ("alcohol", "酒後駕車", ("酒後駕車",), lambda r: r.law_article == "35", ""),
    ("red_light", "闖紅燈", ("闖紅燈",), lambda r: r.law_article == "53" and r.law_paragraph == "1", ""),
    ("speeding", "超速", ("超速",), lambda r: r.law_article in {"40"} or (r.law_article == "33" and r.law_paragraph == "1" and r.law_clause == "1"), ""),
    ("wrong_way", "逆向行駛", ("逆向",), lambda r: r.law_article == "45" and r.law_paragraph == "1" and r.law_clause in {"1", "3"}, ""),
    ("turning", "轉彎未依規定", ("轉彎",), lambda r: r.law_article == "48" and r.law_paragraph == "1", ""),
    ("motorcycle_lane", "機車行駛禁行機車道", ("機車行駛禁行",), lambda r: r.law_article == "45" and r.law_paragraph == "1" and r.law_clause == "13", ""),
    ("two_stage_turn", "機車未依規定兩段式左轉", ("兩段式左轉",), lambda r: r.law_article == "48" and r.law_paragraph == "1" and r.law_clause == "2" and r.law_item == "7", ""),
    ("parking", "併排停車與公車停靠區違規", ("併排", "公車停靠"), lambda r: r.law_article in {"55", "56"}, ""),
    ("large_vehicle", "各式大型車違規", ("大型車",), None, "需提供可辨識的大型車車種與績效管制細項法條。"),
    ("pedestrian", "行人違規", ("行人違規",), lambda r: r.law_article == "78", ""),
    ("yield_pedestrian", "汽機車不禮讓行人", ("不禮讓行人",), lambda r: r.law_article == "44" or (r.law_article == "48" and r.law_paragraph == "2"), ""),
    ("helmet", "未戴安全帽", ("安全帽",), lambda r: r.law_article == "31" and r.law_paragraph == "6", ""),
    ("reckless", "蛇行惡意逼車", ("蛇行",), lambda r: r.law_article == "43" and r.law_paragraph == "1" and r.law_clause in {"1", "3", "4"}, ""),
)


@dataclass(frozen=True)
class PerformanceTargets:
    source_name: str
    units: tuple[str, ...]
    values: dict[str, dict[str, float]]
    matched_headers: dict[str, str]
    warnings: tuple[str, ...] = ()


@dataclass(frozen=True)
class PerformanceStatistics:
    source_names: tuple[str, ...]
    values: dict[str, dict[str, float]]
    available_keys: tuple[str, ...]
    warnings: tuple[str, ...] = ()


def _number(value: Any) -> float | None:
    return float(value) if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def load_performance_targets(path: str | Path, period: str = "month") -> PerformanceTargets:
    """Read the established performance-control workbook's monthly target rows."""
    path = Path(path)
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        return _read_performance_targets(path, workbook, period)
    finally:
        workbook.close()


def _read_performance_targets(path: Path, workbook: Any, period: str) -> PerformanceTargets:
    sheet_name = "13大項月" if period == "month" else "基準值"
    sheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else None
    if sheet is None:
        raise ValueError(f"績效目標值檔找不到「{sheet_name}」工作表。請提供績效管制格式的 Excel。")
    if period == "week":
        headers = {column: str(sheet.cell(2, column).value or "").replace("\n", "") for column in range(2, sheet.max_column + 1)}
        matched_headers = {key: header for key, _label, matchers, _predicate, _requirement in PERFORMANCE_CATEGORIES if (header := next((text for text in headers.values() if any(word in text for word in matchers)), None))}
        values = {}
        for row in range(5, sheet.max_row + 1):
            unit = str(sheet.cell(row, 1).value or "").strip()
            if not unit: continue
            values[unit] = {key: value for key, header in matched_headers.items() if (value := _number(sheet.cell(row, next(column for column, text in headers.items() if text == header)).value)) not in (None, 0)}
        if not values: raise ValueError("週目標值工作表沒有可用的所隊資料。")
        unmatched = [label for key, label, *_ in PERFORMANCE_CATEGORIES if key not in matched_headers]
        return PerformanceTargets(path.name, tuple(values), values, matched_headers, tuple(f"目標檔未對應：{label}" for label in unmatched))
    header_row = target_row = None
    for row in range(1, min(sheet.max_row, 80) + 1):
        if str(sheet.cell(row, 2).value or "").strip() == "目標值":
            target_row, header_row = row, row - 2
            break
    if target_row is None:
        raise ValueError("績效目標值檔找不到「目標值」列。")
    headers = {column: str(sheet.cell(header_row, column).value or "").replace("\n", "") for column in range(3, sheet.max_column + 1)}
    matched_headers = {}
    for key, _label, matchers, _predicate, _requirement in PERFORMANCE_CATEGORIES:
        header = next((text for text in headers.values() if any(word in text for word in matchers)), None)
        if header:
            matched_headers[key] = header
    values: dict[str, dict[str, float]] = {}
    row = target_row
    while row <= sheet.max_row:
        unit = str(sheet.cell(row, 1).value or "").strip()
        marker = str(sheet.cell(row, 2).value or "").strip()
        if unit and marker == "目標值":
            item = {}
            for rule_key, header in matched_headers.items():
                column = next(column for column, text in headers.items() if text == header)
                value = _number(sheet.cell(row, column).value)
                if value is not None and value > 0:
                    item[rule_key] = value
            values[unit] = item
        row += 1
    if not values:
        raise ValueError("績效目標值檔沒有可用的所隊目標值。")
    unmatched = [label for key, label, *_ in PERFORMANCE_CATEGORIES if key not in matched_headers]
    return PerformanceTargets(path.name, tuple(values), values, matched_headers, tuple(f"目標檔未對應：{label}" for label in unmatched))


def _accumulate_performance_statistics(
    workbook: Any,
    values: dict[str, dict[str, float]],
    available: set[str],
) -> None:
    sheet = workbook["13大項月"] if "13大項月" in workbook.sheetnames else workbook.active
    header_row = 4 if sheet.title == "13大項月" else 3
    headers = {column: str(sheet.cell(header_row, column).value or "").replace("\n", "") for column in range(3, sheet.max_column + 1)}
    matches = {key: next((column for column, text in headers.items() if any(word in text for word in matchers)), None) for key, _label, matchers, _predicate, _requirement in PERFORMANCE_CATEGORIES}
    current_unit = ""
    statistic_markers = {"取締件數", "合計"}
    for row in range(header_row + 1, sheet.max_row + 1):
        unit_cell = str(sheet.cell(row, 1).value or "").strip()
        if unit_cell:
            current_unit = unit_cell.replace("板橋分局", "")
        if str(sheet.cell(row, 2).value or "").strip() not in statistic_markers:
            continue
        unit = current_unit
        if not unit or unit == "合計": continue
        for key, column in matches.items():
            if column is None: continue
            value = _number(sheet.cell(row, column).value)
            if value is not None:
                values[unit][key] += value; available.add(key)


def load_performance_statistics(paths: list[str | Path]) -> PerformanceStatistics:
    """Read one or more performance-control workbooks' 取締件數 rows."""
    values: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    available: set[str] = set()
    names, warnings = [], []
    for source_path in paths:
        path = Path(source_path); names.append(path.name)
        # Some official statistical tables have a stale worksheet dimension (A1:U3)
        # despite populated merged rows below; normal mode reads their real extent.
        workbook = load_workbook(path, data_only=True, read_only=False)
        try:
            _accumulate_performance_statistics(workbook, values, available)
        finally:
            workbook.close()
    if not values:
        raise ValueError("統計值檔找不到可用的「取締件數」或「合計」列。請提供績效管制或重大違規項目統計表。")
    return PerformanceStatistics(tuple(names), {unit: dict(item) for unit, item in values.items()}, tuple(sorted(available)), tuple(warnings))


def build_performance(dataset: ViolationDataset | None, targets: PerformanceTargets, start: str | None, end: str | None, statistics: PerformanceStatistics | None = None, selected_keys: list[str] | None = None) -> dict[str, Any]:
    if start and end and start > end:
        raise ValueError("起日不可晚於迄日。")
    selected = set(selected_keys or ())
    source_keys = set(targets.matched_headers) | set(statistics.available_keys if statistics else ())
    categories = [item for item in PERFORMANCE_CATEGORIES if item[0] in source_keys and (not selected or item[0] in selected)]
    unavailable = [
        label for key, label, _matchers, predicate, requirement in categories
        if predicate is None and requirement and not (statistics and key in statistics.available_keys)
    ]
    actual: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    for record in dataset.records if dataset else ():
        if not record.violation_date or (start and record.violation_date < start) or (end and record.violation_date > end):
            continue
        unit = record.enforcement_unit.strip()
        if unit not in targets.values:
            continue
        for key, _label, _matchers, predicate, _requirement in categories:
            if predicate and predicate(record):
                actual[unit][key] += record.count
    rows = []
    totals = {key: {"target": 0.0, "actual": 0.0} for key, *_ in categories}
    for unit in targets.units:
        cells = []
        for key, _label, _matchers, predicate, requirement in categories:
            target = targets.values[unit].get(key, 0.0)
            achieved = statistics.values.get(unit, {}).get(key) if statistics and key in statistics.available_keys else (actual[unit][key] if predicate else None)
            totals[key]["target"] += target
            if achieved is not None: totals[key]["actual"] += achieved
            cells.append({"key": key, "target": target, "actual": achieved, "rate": achieved / target if target and achieved is not None else None, "requirement": requirement})
        rows.append({"unit": unit, "cells": cells})
    warnings = list(targets.warnings)
    missing_targets = [label for key, label, *_ in categories if key not in targets.matched_headers]
    if missing_targets: warnings.append(f"目標值檔缺少：{'、'.join(missing_targets)}。請補上含這些項目的績效管制檔。")
    if unavailable: warnings.append(f"未能計算取締件數：{'、'.join(unavailable)}。{' '.join(requirement for key, _label, _matchers, predicate, requirement in categories if predicate is None)}")
    month_totals = [{"key": key, "target": value["target"], "actual": value["actual"], "difference": value["actual"] - value["target"] if value["actual"] is not None else None, "rate": value["actual"] / value["target"] if value["target"] and value["actual"] is not None else None} for key, value in totals.items()]
    return {"period": {"start": start, "end": end}, "targetSource": targets.source_name, "statisticsSources": list(statistics.source_names) if statistics else [], "warnings": warnings + (list(statistics.warnings) if statistics else []), "categories": [{"key": key, "label": label, "targetHeader": targets.matched_headers.get(key, "未提供"), "available": predicate is not None} for key, label, _matchers, predicate, _requirement in categories], "availableStatisticKeys": list(statistics.available_keys) if statistics else [], "rows": rows, "totals": month_totals}
