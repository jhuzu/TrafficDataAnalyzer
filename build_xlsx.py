#!/usr/bin/env python3
"""Create the standalone accident-analysis workbook with public Python libraries."""

from __future__ import annotations

import json
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


REQUIRED = ("路段", "交叉路名", "發生時間", "肇事原因", "年齡", "當事者區分", "件數")
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
VALUE_BLUE = "DDEBF7"
THIN = Side(style="thin", color="D9E2F3")


def _rank(rows, index, count_index, maximum=20):
    totals = defaultdict(float)
    for row in rows:
        value = str(row[index] or "").strip()
        if value:
            totals[value] += float(row[count_index] or 0)
    return [key for key, _ in sorted(totals.items(), key=lambda item: (-item[1], item[0]))[:maximum]]


def _intersection_rank(rows, road_index, cross_index, count_index):
    totals = defaultdict(float)
    for row in rows:
        road, cross = str(row[road_index] or "").strip(), str(row[cross_index] or "").strip()
        if road and cross:
            totals[f"{road}／{cross}"] += float(row[count_index] or 0)
    return [key for key, _ in sorted(totals.items(), key=lambda item: (-item[1], item[0]))[:20]]


def _header(sheet, left, right, row, title):
    sheet.cell(row, left, title)
    sheet.cell(row, right, "數量")
    for column in range(left, right + 1):
        cell = sheet.cell(row, column)
        cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        cell.font = Font(bold=True)
    sheet.cell(row + 1, left, "依輸入資料自動加總").font = Font(italic=True, color="4B5563", size=10)


def _section(sheet, left, right, title, labels, source_letter, count_letter, last_row, start_row, total_row):
    _header(sheet, left, right, start_row - 2, title)
    for offset, label in enumerate(labels):
        row = start_row + offset
        sheet.cell(row, left, label)
        sheet.cell(row, right, f"=SUMIF('原始資料'!${source_letter}$2:${source_letter}${last_row},{get_column_letter(left)}{row},'原始資料'!${count_letter}$2:${count_letter}${last_row})")
        sheet.cell(row, right).fill = PatternFill("solid", fgColor=VALUE_BLUE)
        sheet.cell(row, right).number_format = "#,##0"
        for column in range(left, right + 1):
            sheet.cell(row, column).border = Border(bottom=THIN)
    sheet.cell(total_row, left, f"{title}合計")
    sheet.cell(total_row, right, f"=SUM('原始資料'!${count_letter}$2:${count_letter}${last_row})")
    for column in range(left, right + 1):
        cell = sheet.cell(total_row, column)
        cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        cell.font = Font(bold=True)
        cell.border = Border(bottom=Side(style="double", color="7F8C8D"))
    sheet.cell(total_row, right).number_format = "#,##0"


def build_workbook(data: dict, output_path: str | Path):
    headers, rows = list(data.get("headers", [])), list(data.get("rows", []))
    missing = [name for name in REQUIRED if name not in headers]
    if missing:
        raise ValueError(f"Pivot Cache 缺少必要欄位：{'、'.join(missing)}")
    index = {name: position for position, name in enumerate(headers)}
    count_index = index["件數"]
    normalized_rows = [list(row) + [""] * max(0, len(headers) - len(row)) for row in rows]
    last_row = len(normalized_rows) + 1
    field = {name: get_column_letter(position + 1) for name, position in index.items()}

    workbook = Workbook()
    source = workbook.active
    source.title = "原始資料"
    source.sheet_view.showGridLines = False
    source.append([*headers, "路口"])
    for row in normalized_rows:
        source.append(row + [None])
    intersection_column = get_column_letter(len(headers) + 1)
    for row in range(2, last_row + 1):
        source.cell(row, len(headers) + 1, f'=IF(AND(${field["路段"]}{row}<>"",${field["交叉路名"]}{row}<>""),${field["路段"]}{row}&"／"&${field["交叉路名"]}{row},"")')
    end_column = get_column_letter(len(headers) + 1)
    table = Table(displayName="RecoveredPivotData", ref=f"A1:{end_column}{last_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False, showColumnStripes=False)
    source.add_table(table)
    for cell in source[1]:
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(bold=True, color="FFFFFF")
    source.freeze_panes = "A2"
    for column in range(1, len(headers) + 2):
        source.column_dimensions[get_column_letter(column)].width = 15
    source.column_dimensions[field["路段"]].width = 22
    source.column_dimensions[field["肇事原因"]].width = 34

    sheet = workbook.create_sheet("分析表")
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:H1")
    sheet["A1"] = "交通事故統計分析"
    sheet["A1"].fill = PatternFill("solid", fgColor=BLUE)
    sheet["A1"].font = Font(bold=True, color="FFFFFF", size=14)
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet.row_dimensions[1].height = 28
    sheet.merge_cells("A2:H2")
    sheet["A2"] = "本表由樞紐快取自動還原；藍底數量為公式加總。"
    sheet["A2"].font = Font(italic=True, color="4B5563", size=10)

    roads = _rank(normalized_rows, index["路段"], count_index)
    intersections = _intersection_rank(normalized_rows, index["路段"], index["交叉路名"], count_index)
    periods = _rank(normalized_rows, index["發生時間"], count_index, 24)
    causes = _rank(normalized_rows, index["肇事原因"], count_index)
    vehicles = _rank(normalized_rows, index["當事者區分"], count_index)
    _section(sheet, 1, 2, "易肇事路段", roads, field["路段"], field["件數"], last_row, 6, 27)
    _section(sheet, 4, 5, "易肇事路口", intersections, intersection_column, field["件數"], last_row, 6, 27)
    _section(sheet, 7, 8, "易肇事時段", periods, field["發生時間"], field["件數"], last_row, 6, 31)
    _section(sheet, 1, 2, "易肇事肇因", causes, field["肇事原因"], field["件數"], last_row, 43, 64)
    _section(sheet, 7, 8, "易肇事車種", vehicles, field["當事者區分"], field["件數"], last_row, 43, 64)
    _header(sheet, 4, 5, 40, "易肇事年齡")
    age_labels = ["0-17", "18-20", "21-30", "31-40", "41-50", "51-60", "61-70", "71歲以上", "沒有年紀", "合計"]
    for offset, label in enumerate(age_labels, 42):
        sheet.cell(offset, 4, label)
    ranges = ((0, 17), (18, 20), (21, 30), (31, 40), (41, 50), (51, 60), (61, 70))
    for row, (low, high) in enumerate(ranges, 42):
        sheet.cell(row, 5, f'=SUMIFS(\'原始資料\'!${field["件數"]}$2:${field["件數"]}${last_row},\'原始資料\'!${field["年齡"]}$2:${field["年齡"]}${last_row},">={low}",\'原始資料\'!${field["年齡"]}$2:${field["年齡"]}${last_row},"<={high}")')
    sheet["E49"] = f'=SUMIFS(\'原始資料\'!${field["件數"]}$2:${field["件數"]}${last_row},\'原始資料\'!${field["年齡"]}$2:${field["年齡"]}${last_row},">=71")'
    sheet["E50"] = f'=SUM(\'原始資料\'!${field["件數"]}$2:${field["件數"]}${last_row})-SUM(E42:E49)'
    sheet["E51"] = "=SUM(E42:E50)"
    for row in range(42, 52):
        for column in range(4, 6):
            sheet.cell(row, column).border = Border(bottom=THIN)
        sheet.cell(row, 5).fill = PatternFill("solid", fgColor=VALUE_BLUE)
        sheet.cell(row, 5).number_format = "#,##0"
    for column in range(4, 6):
        cell = sheet.cell(51, column)
        cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        cell.font = Font(bold=True)
        cell.border = Border(bottom=Side(style="double", color="7F8C8D"))
    for column, width in {"A": 28, "B": 13, "C": 4, "D": 32, "E": 13, "F": 4, "G": 34, "H": 13}.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A5"
    workbook.save(output_path)


def main():
    if len(sys.argv) != 3:
        raise SystemExit("用法：build_xlsx.py 輸入.json 輸出.xlsx")
    build_workbook(json.loads(Path(sys.argv[1]).read_text(encoding="utf-8")), sys.argv[2])


if __name__ == "__main__":
    main()
